import json
import re

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from core.candidate_facts import verifier_formulation


# ============================================================
# CONFIGURATION
# ============================================================

PROFIL_STRUCTURE_PATH = Path(
    "profil_cv_structure.json"
)

PHOTO_PATH = (
    Path("cv")
    / "photo_cv.png"
)

CV_DOCX_VERSION = "cv_docx_v4_recruteur"


# ============================================================
# CHARGEMENT DU PROFIL
# ============================================================

def charger_profil_structure():

    if not PROFIL_STRUCTURE_PATH.exists():

        raise FileNotFoundError(
            "profil_cv_structure.json introuvable."
        )

    return json.loads(
        PROFIL_STRUCTURE_PATH.read_text(
            encoding="utf-8"
        )
    )


# ============================================================
# POLICE
# ============================================================

def regler_police(
    run,
    taille=10,
    gras=False,
    italique=False,
    souligne=False
):

    run.font.name = "Arial"

    run.font.size = Pt(
        taille
    )

    run.font.bold = gras
    run.font.italic = italique
    run.font.underline = souligne

    run._element.get_or_add_rPr()

    run._element.rPr.rFonts.set(
        qn("w:ascii"),
        "Arial"
    )

    run._element.rPr.rFonts.set(
        qn("w:hAnsi"),
        "Arial"
    )


# ============================================================
# PARAGRAPHES
# ============================================================

def regler_paragraphe(
    paragraphe,
    avant=0,
    apres=0,
    interligne=1.0
):

    fmt = paragraphe.paragraph_format

    fmt.space_before = Pt(
        avant
    )

    fmt.space_after = Pt(
        apres
    )

    fmt.line_spacing = interligne


# ============================================================
# BORDURE SOUS LES GRANDES RUBRIQUES
# ============================================================

def ajouter_bordure_bas(
    paragraphe
):

    p = paragraphe._p

    p_pr = p.get_or_add_pPr()

    p_bdr = p_pr.find(
        qn("w:pBdr")
    )

    if p_bdr is None:

        p_bdr = OxmlElement(
            "w:pBdr"
        )

        p_pr.append(
            p_bdr
        )

    bottom = OxmlElement(
        "w:bottom"
    )

    bottom.set(
        qn("w:val"),
        "single"
    )

    bottom.set(
        qn("w:sz"),
        "10"
    )

    bottom.set(
        qn("w:space"),
        "3"
    )

    bottom.set(
        qn("w:color"),
        "000000"
    )

    p_bdr.append(
        bottom
    )


# ============================================================
# SUPPRESSION DES BORDURES DU TABLEAU D'ENTETE
# ============================================================

def supprimer_bordures_table(
    table
):

    tbl = table._tbl
    tbl_pr = tbl.tblPr

    borders = OxmlElement(
        "w:tblBorders"
    )

    for nom in [
        "top",
        "left",
        "bottom",
        "right",
        "insideH",
        "insideV"
    ]:

        bordure = OxmlElement(
            f"w:{nom}"
        )

        bordure.set(
            qn("w:val"),
            "nil"
        )

        borders.append(
            bordure
        )

    tbl_pr.append(
        borders
    )


# ============================================================
# TITRES DE RUBRIQUES
# ============================================================

def ajouter_titre_section(
    document,
    texte
):

    p = document.add_paragraph()

    regler_paragraphe(
        p,
        avant=10,
        apres=5
    )

    # Le titre reste avec le contenu qui suit.
    p.paragraph_format.keep_with_next = True

    run = p.add_run(
        texte.upper()
    )

    regler_police(
        run,
        taille=11.2,
        gras=True
    )

    ajouter_bordure_bas(
        p
    )

    return p


# ============================================================
# DATES
# ============================================================

MOIS_FR = {
    "01": "Janvier",
    "02": "Février",
    "03": "Mars",
    "04": "Avril",
    "05": "Mai",
    "06": "Juin",
    "07": "Juillet",
    "08": "Août",
    "09": "Septembre",
    "10": "Octobre",
    "11": "Novembre",
    "12": "Décembre"
}


def formater_date(
    valeur
):

    valeur = str(
        valeur
        or ""
    ).strip()

    if not valeur:
        return ""

    if valeur.lower() in {
        "aujourd’hui",
        "aujourd'hui"
    }:

        return "Aujourd’hui"

    match = re.fullmatch(
        r"(\d{4})-(\d{2})",
        valeur
    )

    if not match:
        return valeur

    annee = match.group(
        1
    )

    mois = match.group(
        2
    )

    return (
        f"{MOIS_FR.get(mois, mois)} "
        f"{annee}"
    )


def trier_experiences(
    experiences
):

    def cle(
        experience
    ):

        debut = str(
            experience.get(
                "debut",
                "0000-00"
            )
        )

        match = re.match(
            r"(\d{4})(?:-(\d{2}))?",
            debut
        )

        if not match:
            return 0

        annee = int(
            match.group(
                1
            )
        )

        mois = int(
            match.group(
                2
            )
            or 0
        )

        return (
            annee * 100
            + mois
        )

    return sorted(
        experiences,
        key=cle,
        reverse=True
    )


# ============================================================
# ENTETE : NOM / CONTACT / PHOTO
# ============================================================

def ajouter_entete(
    document,
    profil
):

    identite = profil[
        "identite"
    ]

    table = document.add_table(
        rows=1,
        cols=2
    )

    table.autofit = False

    supprimer_bordures_table(
        table
    )

    cellule_gauche = table.cell(
        0,
        0
    )

    cellule_droite = table.cell(
        0,
        1
    )

    cellule_gauche.width = Cm(
        14.8
    )

    cellule_droite.width = Cm(
        3.0
    )

    cellule_gauche.vertical_alignment = (
        WD_CELL_VERTICAL_ALIGNMENT.TOP
    )

    cellule_droite.vertical_alignment = (
        WD_CELL_VERTICAL_ALIGNMENT.TOP
    )

    # --------------------------------------------------------
    # NOM
    # --------------------------------------------------------

    p = cellule_gauche.paragraphs[
        0
    ]

    regler_paragraphe(
        p,
        apres=2
    )

    run = p.add_run(
        identite[
            "nom"
        ]
    )

    regler_police(
        run,
        taille=17,
        gras=True
    )

    # --------------------------------------------------------
    # COORDONNEES
    # --------------------------------------------------------

    coordonnees = [
        (
            f"{identite['ville']} - "
            f"{identite['pays']}"
        ),
        identite[
            "email"
        ],
        identite[
            "telephone"
        ],
        identite[
            "permis"
        ]
    ]

    for ligne in coordonnees:

        p = cellule_gauche.add_paragraph()

        regler_paragraphe(
            p,
            apres=0.5
        )

        run = p.add_run(
            ligne
        )

        regler_police(
            run,
            taille=9.7
        )

    # --------------------------------------------------------
    # PHOTO
    # --------------------------------------------------------

    if PHOTO_PATH.exists():

        p = cellule_droite.paragraphs[
            0
        ]

        p.alignment = (
            WD_ALIGN_PARAGRAPH.RIGHT
        )

        run = p.add_run()

        run.add_picture(
            str(
                PHOTO_PATH
            ),
            width=Cm(
                2.7
            )
        )

    return table


# ============================================================
# TITRE DU CV
# ============================================================

def ajouter_titre_cv(
    document,
    titre
):

    p = document.add_paragraph()

    p.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    regler_paragraphe(
        p,
        avant=4,
        apres=4
    )

    run = p.add_run(
        titre
    )

    regler_police(
        run,
        taille=13.5,
        gras=True
    )


# ============================================================
# PROFIL
# ============================================================

def construire_accroche(
    plan
):

    cv_source = str(
        plan.get(
            "cv_source",
            ""
        )
    ).upper()

    if cv_source == "TECHNICIEN":

        return (
            "Chimiste de formation, orienté synthèse organique "
            "et travail de laboratoire. Titulaire d'un Master 2 "
            "Chimie des Biomolécules et fort de 9 mois de stages "
            "en chimie organique et R&D, je souhaite m'investir "
            "sur un poste de technicien chimiste. À l'aise avec "
            "le suivi réactionnel, la purification, la caractérisation "
            "et le respect des protocoles expérimentaux."
        )

    return (
        "Chimiste titulaire d'un Master 2 Chimie des Biomolécules, "
        "avec 9 mois de stages en synthèse organique et R&D. "
        "Expérience en synthèse multi-étapes, suivi réactionnel, "
        "purification, caractérisation et restitution scientifique."
    )


def ajouter_profil(
    document,
    plan
):

    ajouter_titre_section(
        document,
        "Profil"
    )

    p = document.add_paragraph()

    regler_paragraphe(
        p,
        apres=4,
        interligne=1.08
    )

    run = p.add_run(
        construire_accroche(
            plan
        )
    )

    regler_police(
        run,
        taille=10.0
    )


# ============================================================
# COMPETENCES
# ============================================================

def blocs_competences():

    return [
        (
            "Synthèse organique",
            (
                "Synthèse multi-étapes, suivi de réactions, "
                "préparation de solutions, pesées et dilutions"
            )
        ),

        (
            "Analyse",
            (
                "RMN 1D/2D, IR, UV-Vis, CCM ; "
                "notions en HPLC et LC-MS "
                "(interprétation des résultats)"
            )
        ),

        (
            "Purification",
            (
                "Extraction liquide/liquide, chromatographie "
                "sur colonne, recristallisation"
            )
        ),

        (
            "Laboratoire",
            (
                "Tenue du cahier de laboratoire, rédaction "
                "de rapports scientifiques, respect des protocoles, "
                "travail en équipe"
            )
        ),

        (
            "Logiciels",
            (
                "ChemDraw, MestreNova, SciFinder, Excel"
            )
        )
    ]


def ajouter_competence(
    document,
    titre,
    contenu
):

    p = document.add_paragraph()

    regler_paragraphe(
        p,
        apres=1.8,
        interligne=1.02
    )

    p.paragraph_format.left_indent = Cm(
        0.55
    )

    p.paragraph_format.first_line_indent = Cm(
        -0.30
    )

    run = p.add_run(
        "• "
    )

    regler_police(
        run,
        taille=10
    )

    run = p.add_run(
        f"{titre} : "
    )

    regler_police(
        run,
        taille=10,
        gras=True
    )

    run = p.add_run(
        contenu
    )

    regler_police(
        run,
        taille=10
    )


def ajouter_competences(
    document
):

    ajouter_titre_section(
        document,
        "Compétences techniques"
    )

    for titre, contenu in (
        blocs_competences()
    ):

        ajouter_competence(
            document,
            titre,
            contenu
        )


# ============================================================
# MISSIONS
# ============================================================

def missions_cv(
    experience
):

    poste = str(
        experience.get(
            "poste",
            ""
        )
    ).lower()

    # --------------------------------------------------------
    # M2
    # --------------------------------------------------------

    if "m2" in poste:

        return [
            (
                "Synthèse organique sur des acides aminés "
                "aromatiques cycliques"
            ),

            (
                "Développement et optimisation de "
                "protocoles expérimentaux"
            ),

            (
                "Suivi des réactions et ajustement "
                "des conditions expérimentales"
            ),

            (
                "Purification et caractérisation "
                "des produits obtenus"
            ),

            (
                "Analyses par CCM, RMN et UV-Vis ; "
                "interprétation de résultats HPLC"
            ),

            (
                "Rapports et présentations scientifiques"
            )
        ]

    # --------------------------------------------------------
    # M1
    # --------------------------------------------------------

    if "m1" in poste:

        return [
            (
                "Développement et amélioration "
                "de conditions réactionnelles"
            ),

            (
                "Mise au point de méthodes pour "
                "l'activation d'acides carboxyliques"
            ),

            (
                "Suivi expérimental et ajustements "
                "méthodologiques"
            )
        ]

    return experience.get(
        "missions",
        []
    )


# ============================================================
# PUCES
# ============================================================

def ajouter_puce(
    document,
    texte
):

    p = document.add_paragraph()

    regler_paragraphe(
        p,
        apres=0.8,
        interligne=1.0
    )

    p.paragraph_format.left_indent = Cm(
        0.65
    )

    p.paragraph_format.first_line_indent = Cm(
        -0.30
    )

    run = p.add_run(
        "• "
        + str(
            texte
        )
    )

    regler_police(
        run,
        taille=9.8
    )


# ============================================================
# EXPERIENCE
# ============================================================

def ajouter_experience(
    document,
    experience
):

    debut = formater_date(
        experience.get(
            "debut"
        )
    )

    fin = formater_date(
        experience.get(
            "fin"
        )
    )

    # --------------------------------------------------------
    # DATES
    # --------------------------------------------------------

    p = document.add_paragraph()

    regler_paragraphe(
        p,
        avant=3,
        apres=1
    )

    p.paragraph_format.keep_with_next = True

    run = p.add_run(
        f"{debut} à {fin}"
    )

    regler_police(
        run,
        taille=10,
        souligne=True
    )

    # --------------------------------------------------------
    # POSTE / ORGANISME
    # --------------------------------------------------------

    p = document.add_paragraph()

    regler_paragraphe(
        p,
        apres=2
    )

    p.paragraph_format.keep_with_next = True

    run = p.add_run(
        experience.get(
            "poste",
            ""
        )
    )

    regler_police(
        run,
        taille=10,
        gras=True
    )

    organisation = experience.get(
        "organisation",
        ""
    )

    lieu = experience.get(
        "lieu",
        ""
    )

    if organisation:

        run = p.add_run(
            " - "
            + organisation
        )

        regler_police(
            run,
            taille=9.8,
            italique=True
        )

    if lieu:

        run = p.add_run(
            ", "
            + lieu
        )

        regler_police(
            run,
            taille=9.8,
            italique=True
        )

    # --------------------------------------------------------
    # MISSIONS
    # --------------------------------------------------------

    for mission in missions_cv(
        experience
    ):

        ajouter_puce(
            document,
            mission
        )


# ============================================================
# EXPERIENCES EN LABORATOIRE
# ============================================================

def ajouter_experiences_laboratoire(
    document,
    profil
):

    stages = [
        experience
        for experience
        in profil[
            "experiences"
        ]
        if experience.get(
            "type"
        ) == "stage"
    ]

    stages = trier_experiences(
        stages
    )

    if not stages:
        return

    ajouter_titre_section(
        document,
        "Expériences en laboratoire"
    )

    for experience in stages:

        ajouter_experience(
            document,
            experience
        )


# ============================================================
# EXPERIENCE COMPLEMENTAIRE
# ============================================================

def ajouter_experience_complementaire(
    document,
    profil
):

    experiences = [
        experience
        for experience
        in profil[
            "experiences"
        ]
        if experience.get(
            "type"
        ) != "stage"
    ]

    experiences = trier_experiences(
        experiences
    )

    if not experiences:
        return

    ajouter_titre_section(
        document,
        "Expérience complémentaire"
    )

    for experience in experiences:

        ajouter_experience(
            document,
            experience
        )


# ============================================================
# FORMATION
# ============================================================

def ajouter_formation(
    document,
    formation
):

    p = document.add_paragraph()

    regler_paragraphe(
        p,
        apres=1.8
    )

    p.paragraph_format.left_indent = Cm(
        0.55
    )

    p.paragraph_format.first_line_indent = Cm(
        -0.30
    )

    run = p.add_run(
        "• "
    )

    regler_police(
        run,
        taille=10
    )

    run = p.add_run(
        (
            f"{formation['debut']} - "
            f"{formation['fin']} : "
            f"{formation['diplome']}"
        )
    )

    regler_police(
        run,
        taille=10
    )

    run = p.add_run(
        (
            ", "
            + formation[
                "etablissement"
            ]
        )
    )

    regler_police(
        run,
        taille=10
    )


def ajouter_formations(
    document,
    profil
):

    ajouter_titre_section(
        document,
        "Formation"
    )

    for formation in profil[
        "formation"
    ]:

        ajouter_formation(
            document,
            formation
        )


# ============================================================
# LANGUES
# ============================================================

def ajouter_langues(
    document,
    profil
):

    ajouter_titre_section(
        document,
        "Langues"
    )

    langues = []

    for langue in profil[
        "langues"
    ]:

        texte = (
            f"{langue['langue']} "
            f"({langue['niveau']}"
        )

        details = langue.get(
            "details"
        )

        if details:

            texte += (
                f" - {details}"
            )

        texte += ")"

        langues.append(
            texte
        )

    p = document.add_paragraph()

    regler_paragraphe(
        p,
        apres=0
    )

    p.paragraph_format.left_indent = Cm(
        0.55
    )

    p.paragraph_format.first_line_indent = Cm(
        -0.30
    )

    run = p.add_run(
        "• "
        + " ; ".join(
            langues
        )
    )

    regler_police(
        run,
        taille=10
    )


# ============================================================
# EXTRACTION DU TEXTE DU DOCX
# ============================================================

def texte_document(
    document
):

    textes = []

    for paragraphe in (
        document.paragraphs
    ):

        textes.append(
            paragraphe.text
        )

    for table in document.tables:

        for row in table.rows:

            for cell in row.cells:

                textes.append(
                    cell.text
                )

    return "\n".join(
        textes
    )


# ============================================================
# GARDE-FOU FINAL
# ============================================================

def verifier_cv_final(
    document,
    plan
):

    texte = texte_document(
        document
    )

    controle = verifier_formulation(
        texte
    )

    problemes = list(
        controle.get(
            "problemes",
            []
        )
    )

    texte_normalise = (
        texte.lower()
    )

    # --------------------------------------------------------
    # TERMES INTERDITS / INCONNUS
    # --------------------------------------------------------

    for terme in (
        plan.get(
            "interdits",
            []
        )
        +
        plan.get(
            "inconnus",
            []
        )
    ):

        terme = str(
            terme
            or ""
        ).strip()

        if not terme:
            continue

        if (
            terme.lower()
            in texte_normalise
        ):

            problemes.append(
                {
                    "type":
                        "TERME_NON_AUTORISE_PLAN",

                    "texte":
                        terme
                }
            )

    # --------------------------------------------------------
    # HPLC / LC-MS
    # --------------------------------------------------------

    formulations_interdites = [
        "maîtrise hplc",
        "maitrise hplc",
        "maîtrise de la hplc",
        "maitrise de la hplc",
        "autonome en hplc",
        "expert hplc",

        "maîtrise lc-ms",
        "maitrise lc-ms",
        "maîtrise de la lc-ms",
        "maitrise de la lc-ms",
        "autonome en lc-ms",
        "expert lc-ms",

        "gc-ms"
    ]

    for formulation in (
        formulations_interdites
    ):

        if formulation in texte_normalise:

            problemes.append(
                {
                    "type":
                        "FORMULATION_INTERDITE",

                    "texte":
                        formulation
                }
            )

    if problemes:

        raise RuntimeError(
            "CV refusé par le garde-fou : "
            + json.dumps(
                problemes,
                ensure_ascii=False
            )
        )


# ============================================================
# GENERATION DU CV
# ============================================================

def generer_cv_docx(
    offre,
    plan,
    destination
):

    profil = charger_profil_structure()

    document = Document()

    # --------------------------------------------------------
    # PAGE
    # --------------------------------------------------------

    section = document.sections[
        0
    ]

    section.top_margin = Cm(
        1.15
    )

    section.bottom_margin = Cm(
        1.10
    )

    section.left_margin = Cm(
        1.60
    )

    section.right_margin = Cm(
        1.60
    )

    # --------------------------------------------------------
    # STYLE GENERAL
    # --------------------------------------------------------

    style = document.styles[
        "Normal"
    ]

    style.font.name = "Arial"

    style.font.size = Pt(
        10
    )

    style._element.rPr.rFonts.set(
        qn("w:ascii"),
        "Arial"
    )

    style._element.rPr.rFonts.set(
        qn("w:hAnsi"),
        "Arial"
    )

    # --------------------------------------------------------
    # ENTETE
    # --------------------------------------------------------

    ajouter_entete(
        document,
        profil
    )

    # --------------------------------------------------------
    # TITRE CIBLE
    # --------------------------------------------------------

    titre = plan.get(
        "titre_cv",
        "Technicien chimiste"
    )

    ajouter_titre_cv(
        document,
        titre
    )

    # --------------------------------------------------------
    # PROFIL
    # --------------------------------------------------------

    ajouter_profil(
        document,
        plan
    )

    # --------------------------------------------------------
    # COMPETENCES
    # --------------------------------------------------------

    ajouter_competences(
        document
    )

    # --------------------------------------------------------
    # EXPERIENCES LABORATOIRE
    # --------------------------------------------------------

    ajouter_experiences_laboratoire(
        document,
        profil
    )

    # --------------------------------------------------------
    # EXPERIENCE COMPLEMENTAIRE
    # --------------------------------------------------------

    ajouter_experience_complementaire(
        document,
        profil
    )

    # --------------------------------------------------------
    # FORMATION
    # --------------------------------------------------------

    ajouter_formations(
        document,
        profil
    )

    # --------------------------------------------------------
    # LANGUES
    # --------------------------------------------------------

    ajouter_langues(
        document,
        profil
    )

    # --------------------------------------------------------
    # GARDE-FOU
    # --------------------------------------------------------

    verifier_cv_final(
        document,
        plan
    )

    # --------------------------------------------------------
    # METADONNEES
    # --------------------------------------------------------

    document.core_properties.title = (
        f"CV - {titre}"
    )

    document.core_properties.subject = (
        offre.get(
            "titre",
            ""
        )
    )

    document.core_properties.author = ""

    # --------------------------------------------------------
    # SAUVEGARDE
    # --------------------------------------------------------

    destination = Path(
        destination
    )

    destination.parent.mkdir(
        parents=True,
        exist_ok=True
    )

    document.save(
        destination
    )

    return str(
        destination
    )